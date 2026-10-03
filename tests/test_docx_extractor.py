"""
Unit Test Suite for DOCXExtractor Module
"""

import tempfile
from pathlib import Path
import docx
import pytest

from src.document_processing.docx_extractor import DOCXExtractor
from src.document_processing.exceptions import CorruptedDocumentError, EmptyDocumentError
from src.document_processing.models import ElementType, SourceType


@pytest.fixture
def sample_docx_file():
    """Generates a synthetic DOCX file with headings, paragraphs, and tables."""
    with tempfile.NamedTemporaryFile(suffix=".docx", delete=False) as tmp:
        tmp_path = Path(tmp.name)

    doc = docx.Document()
    doc.add_heading("MASTER SERVICES AGREEMENT", level=1)
    doc.add_paragraph("This Master Services Agreement is entered into as of January 1, 2026.")
    
    doc.add_heading("1. CONFIDENTIALITY", level=2)
    doc.add_paragraph("Each Party agrees to protect Confidential Information using reasonable care.")

    # Add Table
    table = doc.add_table(rows=2, cols=2)
    table.cell(0, 0).text = "Party Name"
    table.cell(0, 1).text = "Role"
    table.cell(1, 0).text = "Acme Corp"
    table.cell(1, 1).text = "Licensor"

    doc.save(tmp_path)
    
    yield tmp_path
    if tmp_path.exists():
        tmp_path.unlink()


@pytest.fixture
def empty_docx_file():
    """Generates an empty DOCX file with no text."""
    with tempfile.NamedTemporaryFile(suffix=".docx", delete=False) as tmp:
        tmp_path = Path(tmp.name)

    doc = docx.Document()
    doc.save(tmp_path)

    yield tmp_path
    if tmp_path.exists():
        tmp_path.unlink()


def test_docx_paragraph_and_heading_extraction(sample_docx_file):
    extractor = DOCXExtractor()
    doc = extractor.extract(sample_docx_file)

    assert doc.source_type == SourceType.DOCX
    assert doc.metadata.paragraph_count > 0
    assert doc.metadata.table_count == 1

    # Check Headings
    headings = [e for e in doc.elements if e.element_type == ElementType.HEADING]
    assert len(headings) == 2
    assert headings[0].heading_level == 1
    assert "MASTER SERVICES AGREEMENT" in headings[0].text
    assert headings[1].heading_level == 2
    assert "CONFIDENTIALITY" in headings[1].text

    # Check Paragraphs
    paragraphs = [e for e in doc.elements if e.element_type == ElementType.PARAGRAPH]
    assert len(paragraphs) >= 2
    assert "January 1, 2026" in paragraphs[0].text


def test_docx_table_extraction(sample_docx_file):
    extractor = DOCXExtractor()
    doc = extractor.extract(sample_docx_file)

    table_cells = [e for e in doc.elements if e.element_type == ElementType.TABLE_CELL]
    assert len(table_cells) == 4
    
    cell_0_0 = [c for c in table_cells if c.row_index == 0 and c.col_index == 0][0]
    assert cell_0_0.text == "Party Name"

    cell_1_0 = [c for c in table_cells if c.row_index == 1 and c.col_index == 0][0]
    assert cell_1_0.text == "Acme Corp"


def test_docx_empty_file(empty_docx_file):
    extractor = DOCXExtractor()
    with pytest.raises(EmptyDocumentError):
        extractor.extract(empty_docx_file)


def test_docx_corrupted_file():
    with tempfile.NamedTemporaryFile(suffix=".docx", delete=False) as tmp:
        tmp.write(b"CORRUPTED_ZIP_BINARY_DATA")
        tmp_path = Path(tmp.name)

    extractor = DOCXExtractor()
    with pytest.raises(CorruptedDocumentError):
        extractor.extract(tmp_path)

    if tmp_path.exists():
        tmp_path.unlink()


def test_docx_missing_file():
    extractor = DOCXExtractor()
    with pytest.raises(FileNotFoundError):
        extractor.extract("non_existent_file_abc999.docx")
