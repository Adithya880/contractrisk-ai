"""
ContractRisk AI - Phase 4 Document Processing Verification & Smoke Test Script

This script verifies end-to-end functionality of document loading, PDF & DOCX extraction,
text normalization, Pydantic schema serialization, and error handling.
"""

import json
import os
import sys
import tempfile
from pathlib import Path

# Add project root to sys.path
sys.path.insert(0, str(Path(__file__).resolve().parent.parent))

import docx
import fitz
from src.document_processing.exceptions import (
    CorruptedDocumentError,
    UnsupportedFileTypeError,
)
from src.document_processing.loader import DocumentLoader
from src.document_processing.models import ExtractedDocument, SourceType


def create_sample_pdf(file_path: Path) -> None:
    """Creates a sample PDF contract file for verification."""
    doc = fitz.open()

    # Page 1
    page1 = doc.new_page()
    page1.insert_text(
        (50, 50),
        "MASTER SOFTWARE LICENSE AGREEMENT\n\n"
        "This Master Software License Agreement (\"Agreement\") is made effective as of "
        "January 15, 2026 (\"Effective Date\"), by and between TechCorp Solutions LLC "
        "(\"Licensor\") and Enterprise Client Inc. (\"Licensee\").\n\n"
        "1. LICENSE GRANT\n"
        "Subject to the terms of this Agreement, Licensor hereby grants to Licensee a non-exclusive, "
        "non-transferable, revocable license to use the Software solely for internal business operations."
    )

    # Page 2
    page2 = doc.new_page()
    page2.insert_text(
        (50, 50),
        "2. LIMITATION OF LIABILITY\n"
        "In no event shall either party be liable for any indirect, incidental, or consequential damages. "
        "Licensor's total cumulative liability shall not exceed $500,000.00 USD.\n\n"
        "3. GOVERNING LAW AND JURISDICTION\n"
        "This Agreement shall be governed by and construed in accordance with the laws of the State of Delaware."
    )

    doc.save(file_path)
    doc.close()


def create_sample_docx(file_path: Path) -> None:
    """Creates a sample DOCX contract file for verification."""
    doc = docx.Document()

    doc.add_heading("NON-DISCLOSURE AND CONFIDENTIALITY AGREEMENT", level=1)
    doc.add_paragraph(
        "This Non-Disclosure Agreement (\"NDA\") is entered into on February 1, 2026, "
        "by and between Global Systems Corp and Partner Innovations Inc."
    )

    doc.add_heading("1. CONFIDENTIAL INFORMATION", level=2)
    doc.add_paragraph(
        "\"Confidential Information\" refers to any proprietary information, technical data, "
        "trade secrets, or know-how disclosed by the Disclosing Party."
    )

    doc.add_heading("2. TERM AND TERMINATION", level=2)
    doc.add_paragraph(
        "This Agreement shall remain in effect for a period of three (3) years from the Effective Date."
    )

    # Add Table
    table = doc.add_table(rows=3, cols=2)
    hdr_cells = table.rows[0].cells
    hdr_cells[0].text = "Party Name"
    hdr_cells[1].text = "Notice Address"

    row1 = table.rows[1].cells
    row1[0].text = "Global Systems Corp"
    row1[1].text = "100 Tech Way, Wilmington, DE 19801"

    row2 = table.rows[2].cells
    row2[0].text = "Partner Innovations Inc"
    row2[1].text = "500 Innovation Blvd, Austin, TX 78701"

    doc.save(file_path)


def run_verification():
    print("=" * 70)
    print(" CONTRACTRISK AI -- PHASE 4 DOCUMENT PROCESSING SMOKE TEST")
    print("=" * 70)

    loader = DocumentLoader()

    with tempfile.TemporaryDirectory() as tmpdir:
        tmpdir_path = Path(tmpdir)
        pdf_path = tmpdir_path / "sample_contract.pdf"
        docx_path = tmpdir_path / "sample_contract.docx"
        txt_path = tmpdir_path / "invalid_file.txt"
        corrupt_pdf_path = tmpdir_path / "corrupt.pdf"

        # Generate test files
        create_sample_pdf(pdf_path)
        create_sample_docx(docx_path)
        txt_path.write_text("Plain text sample.")
        corrupt_pdf_path.write_bytes(b"INVALID_HEADER_DATA_xyz")

        print("\n[TEST 1] Testing PDF Document Extraction...")
        pdf_doc = loader.load_document(pdf_path)
        assert isinstance(pdf_doc, ExtractedDocument)
        assert pdf_doc.source_type == SourceType.PDF
        assert pdf_doc.metadata.page_count == 2
        assert len(pdf_doc.elements) == 2
        assert "LICENSE GRANT" in pdf_doc.full_cleaned_text
        assert "GOVERNING LAW" in pdf_doc.full_cleaned_text
        assert pdf_doc.requires_ocr is False
        print("  [OK] PDF extracted successfully:")
        print(f"       Document ID   : {pdf_doc.document_id}")
        print(f"       Pages         : {pdf_doc.metadata.page_count}")
        print(f"       Element Count : {len(pdf_doc.elements)}")
        print(f"       Cleaned Chars : {len(pdf_doc.full_cleaned_text)}")

        print("\n[TEST 2] Testing DOCX Document Extraction...")
        docx_doc = loader.load_document(docx_path)
        assert isinstance(docx_doc, ExtractedDocument)
        assert docx_doc.source_type == SourceType.DOCX
        assert docx_doc.metadata.paragraph_count > 0
        assert docx_doc.metadata.table_count == 1
        assert "CONFIDENTIAL INFORMATION" in docx_doc.full_cleaned_text
        assert "Global Systems Corp" in docx_doc.full_cleaned_text
        print("  [OK] DOCX extracted successfully:")
        print(f"       Document ID   : {docx_doc.document_id}")
        print(f"       Paragraphs    : {docx_doc.metadata.paragraph_count}")
        print(f"       Tables        : {docx_doc.metadata.table_count}")
        print(f"       Element Count : {len(docx_doc.elements)}")
        print(f"       Cleaned Chars : {len(docx_doc.full_cleaned_text)}")

        print("\n[TEST 3] Verifying Pydantic Model JSON Serialization...")
        pdf_json = pdf_doc.model_dump_json(indent=2)
        parsed_json = json.loads(pdf_json)
        assert parsed_json["source_type"] == "pdf"
        assert len(parsed_json["elements"]) == 2
        print("  [OK] JSON serialization verified.")

        print("\n[TEST 4] Verifying Error Handling & Safety Guards...")

        # 4a. File Not Found
        try:
            loader.load_document(tmpdir_path / "non_existent.pdf")
            print("  [FAIL] Expected FileNotFoundError but none was raised.")
            sys.exit(1)
        except FileNotFoundError:
            print("  [OK] FileNotFoundError correctly raised for missing file.")

        # 4b. Unsupported Extension
        try:
            loader.load_document(txt_path)
            print("  [FAIL] Expected UnsupportedFileTypeError but none was raised.")
            sys.exit(1)
        except UnsupportedFileTypeError:
            print("  [OK] UnsupportedFileTypeError correctly raised for .txt file.")

        # 4c. Corrupted File
        try:
            loader.load_document(corrupt_pdf_path)
            print("  [FAIL] Expected CorruptedDocumentError but none was raised.")
            sys.exit(1)
        except CorruptedDocumentError:
            print("  [OK] CorruptedDocumentError correctly raised for corrupt PDF.")

    print("\n" + "=" * 70)
    print(" [SUCCESS] ALL PHASE 4 SMOKE TESTS PASSED CLEANLY!")
    print("=" * 70)


if __name__ == "__main__":
    run_verification()
