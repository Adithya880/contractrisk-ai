"""
ContractRisk AI - Phase 5 Clause Segmentation Verification & Smoke Test Script

This script verifies end-to-end functionality of loading contracts via Phase 4 DocumentLoader
and segmenting them into structured clauses via Phase 5 ClauseSegmenter.
"""

import json
import os
import sys
import tempfile
from pathlib import Path

# Add project root to sys.path
sys.path.insert(0, str(Path(__file__).resolve().parent.parent))

import docx
import fitz
from src.clause_segmentation import ClauseSegmenter, SegmentationResult
from src.document_processing import DocumentLoader


def create_sample_contract_pdf(file_path: Path) -> None:
    """Creates a synthetic multi-clause PDF contract."""
    doc = fitz.open()

    p1 = doc.new_page()
    p1.insert_text(
        (50, 50),
        "COMMERCIAL LEASE AGREEMENT\n\n"
        "ARTICLE I. DEFINITIONS\n"
        "1.1 Premises. 'Premises' refers to Suite 400 located at 500 Commercial Blvd.\n"
        "1.2 Rent. 'Rent' means the monthly fee of $5,000.00 USD.\n\n"
        "ARTICLE II. INDEMNIFICATION AND LIABILITY\n"
        "2.1 Tenant Indemnity. Tenant shall indemnify Landlord from third-party claims.\n"
        "(a) Tenant liability shall include legal fees.\n"
        "(b) Landlord shall provide prompt written notice."
    )

    doc.save(file_path)
    doc.close()


def create_sample_contract_docx(file_path: Path) -> None:
    """Creates a synthetic multi-clause DOCX contract."""
    doc = docx.Document()

    doc.add_heading("SOFTWARE SERVICES AGREEMENT", level=1)
    doc.add_paragraph("This Software Services Agreement is entered into as of March 1, 2026.")

    doc.add_heading("1. SERVICE LEVEL AGREEMENT", level=2)
    doc.add_paragraph(
        "Provider guarantees 99.9% uptime per calendar month. "
        "Downtime credits shall be calculated as set forth below."
    )

    doc.add_heading("1.1 Downtime Credits", level=3)
    doc.add_paragraph(
        "For uptime between 99.0% and 99.8%, Customer shall receive a 10% credit. "
        "For uptime below 99.0%, Customer shall receive a 25% credit."
    )

    doc.add_heading("2. GOVERNING LAW AND ARBITRATION", level=2)
    doc.add_paragraph("This Agreement shall be governed by the laws of New York State.")

    doc.save(file_path)


def run_verification():
    print("=" * 70)
    print(" CONTRACTRISK AI -- PHASE 5 CLAUSE SEGMENTATION SMOKE TEST")
    print("=" * 70)

    loader = DocumentLoader()
    segmenter = ClauseSegmenter()

    with tempfile.TemporaryDirectory() as tmpdir:
        tmp_path = Path(tmpdir)
        pdf_path = tmp_path / "lease_agreement.pdf"
        docx_path = tmp_path / "services_agreement.docx"

        create_sample_contract_pdf(pdf_path)
        create_sample_contract_docx(docx_path)

        print("\n[TEST 1] Testing Phase 4 -> Phase 5 PDF Clause Segmentation Pipeline...")
        pdf_doc = loader.load_document(pdf_path)
        pdf_result = segmenter.segment(pdf_doc)

        assert isinstance(pdf_result, SegmentationResult)
        assert pdf_result.total_clauses >= 4
        assert pdf_result.text_coverage_ratio > 0.90

        print(f"  [OK] Segmented {pdf_result.total_clauses} clauses from PDF contract:")
        print(f"       Document ID          : {pdf_result.document_id}")
        print(f"       Source Chars         : {pdf_result.source_character_count}")
        print(f"       Segmented Chars      : {pdf_result.segmented_character_count}")
        print(f"       Text Coverage Ratio  : {pdf_result.text_coverage_ratio * 100:.2f}%")

        for c in pdf_result.clauses:
            print(f"       -> Clause {c.clause_id} | Num: {c.clause_number or 'None':<12} | Level: {c.hierarchy_level} | Parent: {c.parent_clause_number or 'None':<6} | Pages: {c.source_pages}")

        print("\n[TEST 2] Testing Phase 4 -> Phase 5 DOCX Clause Segmentation Pipeline...")
        docx_doc = loader.load_document(docx_path)
        docx_result = segmenter.segment(docx_doc)

        assert isinstance(docx_result, SegmentationResult)
        assert docx_result.total_clauses >= 3
        assert docx_result.text_coverage_ratio > 0.90

        print(f"  [OK] Segmented {docx_result.total_clauses} clauses from DOCX contract:")
        print(f"       Document ID          : {docx_result.document_id}")
        print(f"       Source Chars         : {docx_result.source_character_count}")
        print(f"       Segmented Chars      : {docx_result.segmented_character_count}")
        print(f"       Text Coverage Ratio  : {docx_result.text_coverage_ratio * 100:.2f}%")

        for c in docx_result.clauses:
            print(f"       -> Clause {c.clause_id} | Num: {c.clause_number or 'None':<12} | Level: {c.hierarchy_level} | Parent: {c.parent_clause_number or 'None':<6}")

        print("\n[TEST 3] Verifying SegmentationResult Pydantic JSON Serialization...")
        json_output = pdf_result.model_dump_json(indent=2)
        parsed_json = json.loads(json_output)
        assert parsed_json["total_clauses"] == pdf_result.total_clauses
        assert len(parsed_json["clauses"]) == pdf_result.total_clauses
        print("  [OK] JSON serialization verified.")

    print("\n" + "=" * 70)
    print(" [SUCCESS] ALL PHASE 5 SMOKE TESTS PASSED CLEANLY!")
    print("=" * 70)


if __name__ == "__main__":
    run_verification()
